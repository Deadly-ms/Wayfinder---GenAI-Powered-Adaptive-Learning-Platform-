import re
from .ai_engine import ai_engine
import json
import os
from pdfminer.high_level import extract_text
import docx

def extract_text_from_pdf(file_path):
    try:
        return extract_text(file_path)
    except Exception as e:
        print(f"Error reading PDF: {e}")
        return ""

def extract_text_from_docx(file_path):
    try:
        doc = docx.Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        print(f"Error reading DOCX: {e}")
        return ""

def parse_ai_json(response_text):
    try:
        # Determine if we are looking for an Object or a List
        # We pick the first occurrence of '{' or '['
        
        idx_obj = response_text.find('{')
        idx_arr = response_text.find('[')
        
        start_idx = -1
        end_idx = -1
        
        if idx_obj != -1 and (idx_arr == -1 or idx_obj < idx_arr):
            start_idx = idx_obj
            end_idx = response_text.rfind('}') + 1
        elif idx_arr != -1:
            start_idx = idx_arr
            end_idx = response_text.rfind(']') + 1
            
        if start_idx != -1 and end_idx != -1:
            clean_json = response_text[start_idx:end_idx]
            # Remove any potential markdown inside the block (rare but possible)
            clean_json = re.sub(r'```json|```', '', clean_json).strip()
            # Handle standard newlines if they break json
            clean_json = clean_json.replace('\n', ' ')
            
            try:
                data = json.loads(clean_json)
                return data
            except json.JSONDecodeError:
                # Fallback: sometimes single quotes are used
                 try:
                    import ast
                    data = ast.literal_eval(clean_json)
                    return data
                 except:
                    pass

        # Fallback to simple regex if markers not found or parsing failed
        clean_json = re.sub(r'```json|```', '', response_text).strip()
        try:
            data = json.loads(clean_json)
            return data
        except:
             # Last ditch effort
             try:
                import ast
                data = ast.literal_eval(clean_json)
                return data
             except:
                 pass
                 
        return None
    except Exception as e:
        print(f"JSON Parsing Error: {e} \nText: {response_text[:100]}...")
        return None

def analyze_resume(resume_text, target_role, experience_level):
    # Prompt 8: Resume Parser – ATS Scoring Prompt
    prompt_8 = """
    You are an Expert AI Resume Analyzer & Career Coach.

    User Inputs:
    - Target Job Role: {target_role}
    - Experience Level: {experience_level}
    - Resume Content: 
    {resume_text}

    Tasks:
    1. Analyze the resume against the target role.
    2. Calculated ATS compatibility score (0-100). **IMPORTANT: Do NOT round to the nearest 5. Use precise evaluation (e.g., 73, 86, 92).**
    3. Identify 3-4 Key STRENGTHS.
    4. Identify 3-4 Key WEAKNESSES.
    5. List MISSING mandatory skills (be specific).
    6. List actionable skill improvement suggestions. **KEEP SUGGESTIONS CONCISE (max 10 words each).**
    7. Suggest relevant certifications to bridge gaps. **Return ONLY the certification name as a string.**

    Return valid JSON ONLY. 
    IMPORTANT: All lists must be simple arrays of strings. NO OBJECTS.
    {{
        "score": 78, 
        "strengths": ["Strong Python experience", "Good project diversity"],
        "weaknesses": ["Lack of cloud experience", "No leadership examples"],
        "missing_skills": ["Skill A", "Skill B"], 
        "suggestions": ["Learn React hooks", "Build a full-stack project"], 
        "certifications": ["AWS Certified Cloud Practitioner", "Google Data Analytics Certificate"]
    }}
    """
    
    formatted_prompt_8 = prompt_8.format(resume_text=resume_text, target_role=target_role, experience_level=experience_level)
    print("Generating ATS Score...")
    ats_json_str = ai_engine.generate_response(formatted_prompt_8)
    ats_result = parse_ai_json(ats_json_str)
    
    if not ats_result:
        ats_result = {
            "score": 0,
            "strengths": ["Could not parse analysis."],
            "weaknesses": ["Please try again."],
            "missing_skills": ["N/A"],
            "suggestions": ["Retry analysis"],
            "certifications": []
        }

    # Prompt 9: Resume-Based Learning Recommendation Prompt
    prompt_9 = """
    You are an AI Career Learning Advisor.
    Context:
    - Target Role: {target_role}
    - Resume Score: {score}/100
    - Missing Skills: {missing_skills}

    Suggest:
    1. Top 5 specific learning topics/modules to master for this role. **(Simple Strings Only)**
    2. A structured "Learning Path" description (1-2 sentences).
    3. Essential soft skills or extra technical skills required.

    Return valid JSON ONLY.
    IMPORTANT: 'learning_topics' and 'mandatory_skills' MUST be simple arrays of strings. NO OBJECTS.
    {{
        "learning_topics": ["Advanced React", "Node.js Microservices"], 
        "mandatory_skills": ["Communication", "Agile Methodology"], 
        "career_path": "To become a {target_role}, focus on mastering X and Y."
    }}
    """
    
    formatted_prompt_9 = prompt_9.format(
        target_role=target_role, 
        score=ats_result.get('score', 0),
        missing_skills=", ".join(ats_result.get('missing_skills', []))
    )
    
    print("Generating Learning Recommendations...")
    learning_json_str = ai_engine.generate_response(formatted_prompt_9)
    learning_recs = parse_ai_json(learning_json_str)
    
    if not learning_recs:
        learning_recs = {
            "learning_topics": [],
            "mandatory_skills": [],
            "career_path": "Analysis failed to recommend path."
        }
    
    return {
        "ats": ats_result,
        "learning": learning_recs
    }

def generate_better_bullets(bullet_text, target_role):
    prompt = """
    You are an Expert Resume Editor.
    Task: Rewrite the following resume bullet point to make it "Harvard-style" (Action-Result oriented).
    Target Role: {target_role}
    
    Weak Bullet: "{bullet_text}"
    
    Generate 3 distinct, powerful variations:
    1. Quantified results focused.
    2. Leadership/Action focused.
    3. Technical/Skill focused.

    Keep variations concise and impactful (max 15 words).
    
    Return as JSON only: ["Variation 1...", "Variation 2...", "Variation 3..."]
    """
    try:
        response = ai_engine.generate_response(prompt.format(bullet_text=bullet_text, target_role=target_role))
        return parse_ai_json(response) or ["Could not generate variations."]
    except:
        return ["Error generating variations."]

def generate_cover_letter(resume_text, job_description):
    prompt = """
    You are an Expert Career Coach.
    Task: Write a compelling, highly personalized Cover Letter.
    
    My Resume Summary:
    {resume_text}
    
    Target Job Description:
    {job_description}
    
    Guidelines:
    - distinct "Hook" opening.
    - Connect specific past projects to JD requirements.
    - Professional but enthusiastic tone.
    - KEEP IT CONCISE (max 300 words).
    
    Return the cover letter text directly (No JSON, no markdown wrappers).
    """
    short_resume = resume_text[:3000]
    response = ai_engine.generate_response(prompt.format(resume_text=short_resume, job_description=job_description))
    return response.strip()

def generate_rewrite_suggestions(resume_text, target_role):
    prompt = """
    You are an Expert Resume Editor (Harvard/McKinsey style).
    Task: Analyze the resume and identify 3-4 specific "Weak Areas" (e.g., a weak Profile Summary, or a vague Experience bullet point).
    Target Role: {target_role}
    
    Resume Content:
    {resume_text}
    
    For EACH weak area found:
    1. Quote the "Original" text.
    2. Explain "Why it's weak" (1 sentence).
    Generate 3 "Harvard-style" Variations (Action + Result + Quantified). 
    **KEEP VARIATIONS CONCISE (max 20 words).**
    **DO NOT USE MARKDOWN formatting (no bold **, no italics *). Plain text only.**
    
    Return as JSON List:
    [
        {{
            "section": "Profile Summary" or "Experience: Company Name",
            "original": "...",
            "issue": "...",
            "variations": ["Var 1", "Var 2", "Var 3"]
        }},
        ...
    ]
    """
    try:
        short_resume = resume_text[:4000]
        response = ai_engine.generate_response(prompt.format(resume_text=short_resume, target_role=target_role))
        return parse_ai_json(response) or []
    except:
        return []

def create_cover_letter_docx(text, filename="Cover_Letter.docx"):
    doc = docx.Document()
    
    # Simple Style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = docx.shared.Pt(11)
    
    # Add content (handling newlines)
    for line in text.split('\n'):
        doc.add_paragraph(line)
        
    save_path = os.path.join("instance", filename)
    doc.save(save_path)
    return save_path
